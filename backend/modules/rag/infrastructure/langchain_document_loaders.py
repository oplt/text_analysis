from __future__ import annotations

import asyncio
import logging
import tempfile
import unicodedata
from pathlib import Path

from backend.core.config import settings
from backend.modules.rag.domain.models import ParsedDocument

logger = logging.getLogger(__name__)


def _decode_text(content: bytes) -> tuple[str, dict[str, int | str]]:
    """Decode and normalize text while surfacing lossy input deterministically."""
    try:
        decoded = content.decode("utf-8-sig")
        replacement_count = 0
    except UnicodeDecodeError:
        decoded = content.decode("utf-8", errors="replace")
        replacement_count = decoded.count("\ufffd")
    if (
        decoded
        and replacement_count / len(decoded) > settings.RAG_TEXT_DECODE_MAX_REPLACEMENT_RATIO
    ):
        raise ValueError("Text decoding corruption exceeds the supported replacement threshold")
    return unicodedata.normalize("NFC", decoded), {
        "text_encoding": "utf-8",
        "unicode_normalization": "NFC",
        "normalization_version": "unicode-nfc-v1",
        "decode_replacement_count": replacement_count,
    }


def _parse_text(content: bytes) -> list[ParsedDocument]:
    text, metadata = _decode_text(content)
    text = text.strip()
    if not text:
        return []
    return [ParsedDocument(content=text, metadata={"format": "text", **metadata})]


def _parse_markdown(content: bytes) -> list[ParsedDocument]:
    text, decode_metadata = _decode_text(content)
    text = text.strip()
    if not text:
        return []
    sections: list[ParsedDocument] = []
    heading: str | None = None
    body: list[str] = []

    def flush() -> None:
        section = "\n".join(body).strip()
        if section:
            sections.append(
                ParsedDocument(
                    content=section,
                    metadata={
                        "format": "markdown",
                        "section_heading": heading,
                        **decode_metadata,
                    },
                )
            )

    for line in text.splitlines():
        if line.lstrip().startswith("#"):
            flush()
            body = [line]
            heading = line.lstrip("#").strip() or None
        else:
            body.append(line)
    flush()
    return sections


def _looks_numeric(value: str) -> bool:
    candidate = value.strip().replace(",", "")
    if not candidate:
        return False
    try:
        float(candidate)
        return True
    except ValueError:
        return False


def _parse_csv(content: bytes) -> list[ParsedDocument]:
    """Row/record-aware CSV parsing — never flatten the entire file into one blob."""
    import csv
    import io

    text, decode_metadata = _decode_text(content)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []

    headers = [cell.strip() for cell in rows[0]]
    try:
        has_header = csv.Sniffer().has_header(text)
    except csv.Error:
        has_header = False
    # Prefer headerless when the first row is entirely numeric-like; Sniffer
    # otherwise invents headers for short synthetic fixtures.
    if has_header and headers and all(_looks_numeric(cell) for cell in headers):
        has_header = False
    data_rows = rows[1:] if has_header else rows
    docs: list[ParsedDocument] = []
    for row_number, row in enumerate(data_rows, start=1 if has_header else 0):
        cells = [cell.strip() for cell in row]
        if not any(cells):
            continue
        if has_header:
            pairs = [
                f"{headers[i] if i < len(headers) else f'col_{i}'}: {cells[i]}"
                for i in range(len(cells))
                if cells[i]
            ]
            content_text = " | ".join(pairs)
        else:
            content_text = " | ".join(c for c in cells if c)
        docs.append(
            ParsedDocument(
                content=content_text,
                metadata={
                    "format": "csv",
                    "row_number": row_number,
                    "headers": headers if has_header else None,
                    "column_count": len(cells),
                    "parser": "csv-row-v1",
                    "header_detected": has_header,
                    **decode_metadata,
                },
            )
        )
    return docs


def _parse_pdf(content: bytes) -> list[ParsedDocument]:
    from backend.modules.rag.infrastructure.pdf_parsers import parse_pdf_with_fallback

    return parse_pdf_with_fallback(content)


def _parse_docx(content: bytes) -> list[ParsedDocument]:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for DOCX parsing. Install with: uv add python-docx"
        ) from exc

    import io

    document = docx.Document(io.BytesIO(content))
    sections: list[ParsedDocument] = []
    heading: str | None = None
    paragraphs: list[str] = []

    def flush() -> None:
        text = "\n".join(paragraphs).strip()
        if text:
            sections.append(
                ParsedDocument(
                    content=text,
                    metadata={
                        "format": "docx",
                        "section_heading": heading,
                        "unicode_normalization": "NFC",
                        "normalization_version": "unicode-nfc-v1",
                        "parser": "docx-section-v1",
                    },
                )
            )

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = getattr(paragraph.style, "name", "") or ""
        if style_name.lower().startswith("heading"):
            flush()
            paragraphs = [text]
            heading = text
        else:
            paragraphs.append(text)
    flush()
    for table_index, table in enumerate(document.tables):
        rows = [
            [unicodedata.normalize("NFC", cell.text).strip() for cell in row.cells]
            for row in table.rows
        ]
        if any(any(cell for cell in row) for row in rows):
            sections.append(
                ParsedDocument(
                    content="\n".join(" | ".join(row) for row in rows),
                    metadata={
                        "format": "docx",
                        "parsed_block_type": "table",
                        "table_index": table_index,
                        "section_heading": heading,
                        "unicode_normalization": "NFC",
                        "normalization_version": "unicode-nfc-v1",
                        "parser": "docx-table-v1",
                    },
                )
            )
    return sections


def _parse_with_langchain_loader(file_path: str, loader_name: str) -> list[ParsedDocument]:
    """Fallback LangChain community loaders when available."""
    try:
        if loader_name == "text":
            from langchain_community.document_loaders import TextLoader

            loader = TextLoader(file_path, encoding="utf-8")
        elif loader_name == "csv":
            from langchain_community.document_loaders import CSVLoader

            loader = CSVLoader(file_path)
        else:
            return []
        lc_docs = loader.load()
        return [
            ParsedDocument(
                content=doc.page_content,
                metadata=dict(doc.metadata),
                page_number=doc.metadata.get("page"),
            )
            for doc in lc_docs
            if doc.page_content.strip()
        ]
    except Exception:
        logger.debug("LangChain loader %s unavailable, using native parser", loader_name)
        return []


def _parse_bytes_sync(
    *,
    content: bytes,
    filename: str,
    content_type: str,
) -> list[ParsedDocument]:
    """Parse raw bytes into documents. CPU-bound — use ``load_documents_from_bytes``."""
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix in {"md", "markdown"}:
        return _parse_markdown(content)
    if suffix == "txt" or content_type.startswith("text/"):
        return _parse_text(content)
    if suffix == "csv":
        parsed = _parse_csv(content)
        if parsed:
            return parsed
    if suffix == "pdf" or content_type == "application/pdf":
        return _parse_pdf(content)
    if suffix == "docx" or content_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }:
        return _parse_docx(content)

    with tempfile.NamedTemporaryFile(suffix=f".{suffix}", delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name
    try:
        return _parse_with_langchain_loader(tmp_path, suffix)
    finally:
        Path(tmp_path).unlink(missing_ok=True)


async def load_documents_from_bytes(
    *,
    content: bytes,
    filename: str,
    content_type: str,
) -> list[ParsedDocument]:
    """Parse file bytes into ParsedDocument list off the event loop thread."""
    return await asyncio.to_thread(
        _parse_bytes_sync,
        content=content,
        filename=filename,
        content_type=content_type,
    )
