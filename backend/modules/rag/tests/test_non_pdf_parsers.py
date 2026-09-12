from __future__ import annotations

import unittest

from backend.modules.rag.infrastructure.langchain_document_loaders import (
    _parse_csv,
    _parse_docx,
    _parse_text,
)


class NonPdfParserTests(unittest.TestCase):
    def test_csv_header_and_headerless_rows(self):
        with_header = _parse_csv(b"name,score\nalice,1\nbob,2\n")
        self.assertTrue(with_header)
        self.assertTrue(with_header[0].metadata["header_detected"])
        self.assertIn("name:", with_header[0].content)
        self.assertEqual(with_header[0].metadata["unicode_normalization"], "NFC")
        self.assertEqual(with_header[0].metadata["normalization_version"], "unicode-nfc-v1")

        headerless = _parse_csv(b"10,20,30\n40,50,60\n70,80,90\n")
        self.assertTrue(headerless)
        self.assertFalse(headerless[0].metadata["header_detected"])
        self.assertIsNone(headerless[0].metadata["headers"])
        self.assertEqual(headerless[0].metadata["normalization_version"], "unicode-nfc-v1")
        self.assertIn(" | ", headerless[0].content)

    def test_text_nfc_normalization_provenance(self):
        # cafe + combining acute accent should normalize to composed form
        docs = _parse_text("cafe\u0301".encode())
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].content, "caf\u00e9")
        self.assertEqual(docs[0].metadata["unicode_normalization"], "NFC")
        self.assertEqual(docs[0].metadata["normalization_version"], "unicode-nfc-v1")

    def test_docx_tables_are_emitted_as_atomic_blocks(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx unavailable")

        from io import BytesIO

        document = docx.Document()
        document.add_paragraph("Intro")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"
        buffer = BytesIO()
        document.save(buffer)
        docs = _parse_docx(buffer.getvalue())
        table_docs = [doc for doc in docs if doc.metadata.get("parsed_block_type") == "table"]
        self.assertEqual(len(table_docs), 1)
        self.assertIn("A | B", table_docs[0].content)
        self.assertEqual(table_docs[0].metadata["normalization_version"], "unicode-nfc-v1")
